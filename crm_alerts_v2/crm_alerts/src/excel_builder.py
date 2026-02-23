"""
Excel Report Builder
=====================
Sheets: Owner Summary | Master List | Leads No Action | Leads Not Converted | Per-Stage sheets
Also:   Per-owner Excel files for individual alerts
"""

import logging
from datetime import date, datetime
from collections import defaultdict
from pathlib import Path
from .analyzer import get_field_value, resolve_owner, get_kva_category, business_days_between

log = logging.getLogger(__name__)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ================================================================
# STYLE HELPERS
# ================================================================

def _styles() -> dict:
    return {
        "hf":         Font(bold=True, color="FFFFFF", name="Arial", size=11),
        "hfill":      PatternFill("solid", fgColor="2C3E50"),
        "ha":         Alignment(horizontal="center", vertical="center", wrap_text=True),
        "df":         Font(name="Arial", size=10),
        "rf":         Font(name="Arial", size=10, bold=True, color="C0392B"),
        "bf":         Font(name="Arial", size=10, bold=True),
        "b":          Border(
                          left=Side("thin", "CCCCCC"),
                          right=Side("thin", "CCCCCC"),
                          top=Side("thin", "CCCCCC"),
                          bottom=Side("thin", "CCCCCC"),
                      ),
        "redfill":    PatternFill("solid", fgColor="FDECEA"),
        "yellowfill": PatternFill("solid", fgColor="FFF8E1"),
    }


def _hdr(ws, headers: list, s: dict):
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = s["hf"]
        c.fill = s["hfill"]
        c.alignment = s["ha"]
        c.border = s["b"]


def _autowidth(ws, headers: list, row_count: int):
    for ci in range(1, len(headers) + 1):
        col = ws.cell(row=1, column=ci).column_letter
        max_len = len(str(headers[ci - 1]))
        for ri in range(2, min(row_count + 2, 102)):
            val = ws.cell(row=ri, column=ci).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col].width = min(max_len + 4, 50)


def _filter(ws, headers: list, row_count: int):
    if row_count > 0:
        end = ws.cell(row=1, column=len(headers)).column_letter
        ws.auto_filter.ref = f"A1:{end}{row_count + 1}"


def _calc_age(record: dict) -> object:
    """Return calendar days since record was created, or '-' if unavailable."""
    created = record.get("Created_Time", "")
    if not created or str(created) in ("null", "None"):
        return "-"
    try:
        created_date = datetime.strptime(str(created)[:10], "%Y-%m-%d").date()
        return (date.today() - created_date).days
    except Exception:
        return "-"


# ================================================================
# MAIN REPORT (for MD / Management)
# ================================================================

def build_excel_report(overdue_no_action, overdue_not_converted,
                       stage_results, config, user_cache, crm) -> str:
    if not HAS_OPENPYXL:
        log.warning("openpyxl not installed — cannot generate Excel.")
        return None

    wb = Workbook()
    s = _styles()
    stage_names = config["stage_display_names"]
    stages = [st for st in config["stages_to_monitor"] if st not in config["terminal_stages"]]

    # ----------------------------------------------------------------
    # SHEET 1: OWNER SUMMARY
    # ----------------------------------------------------------------
    log.info("  Building Owner Summary...")
    ws = wb.active
    ws.title = "Owner Summary"
    ws.sheet_properties.tabColor = "2C3E50"

    owner_data = defaultdict(lambda: {
        "leads_na": 0, "leads_nc": 0, "total": 0, "amount": 0,
        "days_over": 0, "stages": defaultdict(int),
    })

    for item in overdue_no_action:
        n, _ = resolve_owner(item["lead"], user_cache)
        owner_data[n]["leads_na"] += 1

    for item in overdue_not_converted:
        n, _ = resolve_owner(item["lead"], user_cache)
        owner_data[n]["leads_nc"] += 1

    for stage, items in stage_results.items():
        for item in items:
            n, _ = resolve_owner(item["record"], user_cache)
            od = owner_data[n]
            od["total"] += 1
            od["stages"][stage] += 1
            if item["threshold"] > 0:
                od["days_over"] += max(0, item["days_stalled"] - item["threshold"])
            amt = item["record"].get("Amount")
            if amt and str(amt) not in ("null", "None"):
                try:
                    od["amount"] += float(amt)
                except (ValueError, TypeError):
                    pass

    hdrs = ["Owner", "Leads (No Action)", "Leads (Not Conv.)", "Total Enq Stalled"]
    for st in stages:
        hdrs.append(stage_names.get(st, st))
    hdrs.extend(["Total Amount (₹)", "Total Days Stalled For", "Avg Days Stalled For"])
    _hdr(ws, hdrs, s)

    sorted_owners = sorted(owner_data.items(), key=lambda x: x[1]["total"], reverse=True)
    for ri, (name, od) in enumerate(sorted_owners, 2):
        row = [name, od["leads_na"], od["leads_nc"], od["total"]]
        for st in stages:
            row.append(od["stages"].get(st, 0))
        avg = round(od["days_over"] / od["total"], 1) if od["total"] > 0 else 0
        row.extend([od["amount"], od["days_over"], avg])

        for ci, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.border = s["b"]
            c.font = s["df"]
            if ci == 4 and val > 20:
                c.font = s["rf"]
                c.fill = s["redfill"]
            if ci == len(row) and val > 5:
                c.font = s["rf"]
                c.fill = s["redfill"]

        ws.cell(row=ri, column=len(hdrs) - 2).number_format = "#,##0"

    _autowidth(ws, hdrs, len(sorted_owners))
    _filter(ws, hdrs, len(sorted_owners))
    ws.freeze_panes = "A2"

    # ----------------------------------------------------------------
    # SHEET 2: MASTER LIST
    # ----------------------------------------------------------------
    log.info("  Building Master List...")
    ws2 = wb.create_sheet("Master List")
    ws2.sheet_properties.tabColor = "E74C3C"

    m_hdrs = [
        "Owner", "Customer Name", "Enquiry Name", "Enq No.", "Stage",
        "kVA", "Offering", "Amount (₹)", "Threshold",
        "Stalled For", "Enquiry Age (Days)", "Created Date", "Last Updated",
        "Stage Since", "Remarks",
    ]
    _hdr(ws2, m_hdrs, s)

    master = []
    for stage, items in stage_results.items():
        for item in items:
            rec = item["record"]
            owner_name, _ = resolve_owner(rec, user_cache)
            is_na = item["threshold"] == 0
            stalled_for = "NA" if is_na else max(0, item["days_stalled"] - item["threshold"])
            threshold_display = "NA" if is_na else item["threshold"]
            age = _calc_age(rec)
            created = get_field_value(rec, "Created_Time", user_cache, crm)
            updated = get_field_value(rec, "Modified_Time", user_cache, crm)

            amt = None
            raw_amt = rec.get("Amount")
            if raw_amt and str(raw_amt) not in ("null", "None"):
                try:
                    amt = float(raw_amt)
                except (ValueError, TypeError):
                    pass

            master.append({
                "row": [
                    owner_name,
                    get_field_value(rec, "Account_Name", user_cache, crm),
                    get_field_value(rec, "Deal_Name", user_cache, crm),
                    get_field_value(rec, "ENQ_NUM", user_cache, crm),
                    stage_names.get(stage, stage),
                    get_field_value(rec, "DG_KVA", user_cache, crm),
                    get_field_value(rec, "OFFERING", user_cache, crm),
                    amt,
                    threshold_display,
                    stalled_for,
                    age,
                    created[:10] if created != "-" and len(created) >= 10 else created,
                    updated[:10] if updated != "-" and len(updated) >= 10 else updated,
                    updated[:10] if updated != "-" and len(updated) >= 10 else updated,
                    get_field_value(rec, "Status_Remarks", user_cache, crm),
                ],
                "stalled_for": stalled_for if isinstance(stalled_for, int) else 0,
            })

    master.sort(key=lambda x: (x["row"][0], -x["stalled_for"]))

    for ri, item in enumerate(master, 2):
        row = item["row"]
        for ci, val in enumerate(row, 1):
            c = ws2.cell(row=ri, column=ci, value=val)
            c.border = s["b"]
            c.font = s["df"]

        sf_val = item["stalled_for"]
        sf = ws2.cell(row=ri, column=10)
        if isinstance(sf_val, int):
            if sf_val >= 10:
                sf.font = s["rf"]
                sf.fill = s["redfill"]
            elif sf_val >= 5:
                sf.font = s["bf"]
                sf.fill = s["yellowfill"]

        if row[7] is not None:
            ws2.cell(row=ri, column=8).number_format = "#,##0"

    _autowidth(ws2, m_hdrs, len(master))
    _filter(ws2, m_hdrs, len(master))
    ws2.freeze_panes = "A2"

    # ----------------------------------------------------------------
    # SHEETS 3 & 4: LEADS
    # ----------------------------------------------------------------
    log.info("  Building Lead sheets...")
    l_hdrs = (
        [c["label"] for c in config["lead_columns"]]
        + ["Created Date", "Lead Age (Days)", "Biz Days Overdue"]
    )

    def _lead_sheet(ws_l, items, color):
        ws_l.sheet_properties.tabColor = color
        _hdr(ws_l, l_hdrs, s)
        for ri, item in enumerate(items, 2):
            row = [
                get_field_value(item["lead"], c["field"], user_cache, crm)
                for c in config["lead_columns"]
            ]
            created = get_field_value(item["lead"], "Created_Time", user_cache, crm)
            created_short = created[:10] if created != "-" and len(created) >= 10 else created
            age = _calc_age(item["lead"])
            row.extend([created_short, age, item["biz_days"]])
            for ci, val in enumerate(row, 1):
                c = ws_l.cell(row=ri, column=ci, value=val)
                c.border = s["b"]
                c.font = s["rf"] if ci == len(row) else s["df"]
        _autowidth(ws_l, l_hdrs, len(items))
        _filter(ws_l, l_hdrs, len(items))
        ws_l.freeze_panes = "A2"

    _lead_sheet(wb.create_sheet("Leads - No Action"), overdue_no_action, "C0392B")
    _lead_sheet(wb.create_sheet("Leads - Not Converted"), overdue_not_converted, "E67E22")

    # ----------------------------------------------------------------
    # SHEETS 5+: PER STAGE
    # ----------------------------------------------------------------
    log.info("  Building per-stage sheets...")
    e_hdrs = (
        [c["label"] for c in config["enquiry_columns"]]
        + ["kVA Category", "Threshold", "Stalled For", "Enquiry Age (Days)", "Stage Since"]
    )

    for stage in stages:
        items = stage_results.get(stage, [])
        if not items:
            continue

        display = stage_names.get(stage, stage)
        ws_stg = wb.create_sheet(display[:31])
        ws_stg.sheet_properties.tabColor = (
            config["stage_colors"].get(stage, "2c3e50").replace("#", "")
        )
        _hdr(ws_stg, e_hdrs, s)

        for ri, item in enumerate(items, 2):
            rec = item["record"]
            is_na = item["threshold"] == 0
            stalled_for = "NA" if is_na else max(0, item["days_stalled"] - item["threshold"])
            threshold_display = "NA" if is_na else f"{item['threshold']} day(s)"
            age = _calc_age(rec)
            mod = get_field_value(rec, "Modified_Time", user_cache, crm)
            stage_since = mod[:10] if mod != "-" and len(mod) >= 10 else mod

            row = [
                get_field_value(rec, c["field"], user_cache, crm)
                for c in config["enquiry_columns"]
            ]
            row.extend([item["kva_category"], threshold_display, stalled_for, age, stage_since])

            for ci, val in enumerate(row, 1):
                c = ws_stg.cell(row=ri, column=ci, value=val)
                c.border = s["b"]
                c.font = s["df"]

            sf_ci = len(row) - 2
            sf = ws_stg.cell(row=ri, column=sf_ci)
            if isinstance(stalled_for, int):
                if stalled_for >= 10:
                    sf.font = s["rf"]
                    sf.fill = s["redfill"]
                elif stalled_for >= 5:
                    sf.font = s["bf"]
                    sf.fill = s["yellowfill"]

        _autowidth(ws_stg, e_hdrs, len(items))
        _filter(ws_stg, e_hdrs, len(items))
        ws_stg.freeze_panes = "A2"

    # ----------------------------------------------------------------
    # SAVE
    # ----------------------------------------------------------------
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    filename = Path(__file__).parent.parent / f"stalled_report_{timestamp}.xlsx"
    wb.save(filename)
    log.info(f"  Excel saved: {filename}")
    return str(filename)


# ================================================================
# OWNER EXCEL (individual report per owner)
# ================================================================

def build_owner_excel(owner_email, owner_name, lead_alerts, enq_alerts,
                      config, user_cache, crm) -> str:
    """Build a personalised Excel for a single owner containing only their items."""
    if not HAS_OPENPYXL:
        return None

    wb = Workbook()
    s = _styles()
    stage_names = config["stage_display_names"]

    # ----------------------------------------------------------------
    # SHEET 1: SUMMARY
    # ----------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_properties.tabColor = "2C3E50"

    ws["A1"] = f"Stalled Items Report — {owner_name}"
    ws["A1"].font = Font(bold=True, size=14, name="Arial", color="2C3E50")
    ws["A2"] = f"Date: {date.today().strftime('%d-%b-%Y, %A')}"
    ws["A2"].font = Font(size=11, name="Arial")
    ws["A3"] = "Please review and take action on the items below."
    ws["A3"].font = Font(size=10, name="Arial", color="C0392B")

    ws["A5"] = "Section"
    ws["B5"] = "Count"
    ws["A5"].font = s["hf"]
    ws["A5"].fill = s["hfill"]
    ws["A5"].border = s["b"]
    ws["B5"].font = s["hf"]
    ws["B5"].fill = s["hfill"]
    ws["B5"].border = s["b"]

    lead_na_count = len([a for a in lead_alerts if a.get("type") == "No Action"])
    lead_nc_count = len([a for a in lead_alerts if a.get("type") == "Not Converted"])

    stage_counts = defaultdict(int)
    for ea in enq_alerts:
        stage_counts[ea.get("stage", "Unknown")] += 1

    row_i = 6

    ws.cell(row=row_i, column=1, value="Leads — No Action").font = s["df"]
    ws.cell(row=row_i, column=1).border = s["b"]
    c = ws.cell(row=row_i, column=2, value=lead_na_count)
    c.font = s["rf"] if lead_na_count > 0 else s["df"]
    c.border = s["b"]
    row_i += 1

    ws.cell(row=row_i, column=1, value="Leads — Not Converted").font = s["df"]
    ws.cell(row=row_i, column=1).border = s["b"]
    c = ws.cell(row=row_i, column=2, value=lead_nc_count)
    c.font = s["rf"] if lead_nc_count > 0 else s["df"]
    c.border = s["b"]
    row_i += 1

    for stage in config["stages_to_monitor"]:
        if stage in config["terminal_stages"]:
            continue
        cnt = stage_counts.get(stage, 0)
        if cnt == 0:
            continue
        display = stage_names.get(stage, stage)
        ws.cell(row=row_i, column=1, value=f"Enquiries — {display}").font = s["df"]
        ws.cell(row=row_i, column=1).border = s["b"]
        c = ws.cell(row=row_i, column=2, value=cnt)
        c.font = s["rf"] if cnt > 0 else s["df"]
        c.border = s["b"]
        row_i += 1

    total = lead_na_count + lead_nc_count + len(enq_alerts)
    ws.cell(row=row_i, column=1, value="TOTAL").font = s["bf"]
    ws.cell(row=row_i, column=1).border = s["b"]
    ws.cell(row=row_i, column=2, value=total).font = s["rf"]
    ws.cell(row=row_i, column=2).border = s["b"]

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 12

    # ----------------------------------------------------------------
    # SHEET 2: MY LEADS
    # ----------------------------------------------------------------
    ws_leads = wb.create_sheet("My Leads")
    ws_leads.sheet_properties.tabColor = "C0392B"

    l_hdrs = (
        [c["label"] for c in config["lead_columns"]]
        + ["Issue", "Created Date", "Lead Age (Days)", "Biz Days Overdue"]
    )
    _hdr(ws_leads, l_hdrs, s)

    if lead_alerts:
        for ri, item in enumerate(lead_alerts, 2):
            ld = item["lead"]
            row = [
                get_field_value(ld, c["field"], user_cache, crm)
                for c in config["lead_columns"]
            ]
            created = get_field_value(ld, "Created_Time", user_cache, crm)
            created_short = created[:10] if created != "-" and len(created) >= 10 else created
            age = _calc_age(ld)
            row.extend([item["type"], created_short, age, item["days"]])
            for ci, val in enumerate(row, 1):
                c = ws_leads.cell(row=ri, column=ci, value=val)
                c.border = s["b"]
                c.font = s["df"]
            ws_leads.cell(row=ri, column=len(row) - 3).font = s["rf"]
            ws_leads.cell(row=ri, column=len(row)).font = s["rf"]
        _autowidth(ws_leads, l_hdrs, len(lead_alerts))
        _filter(ws_leads, l_hdrs, len(lead_alerts))
    else:
        ws_leads["A3"] = "No stalled leads assigned to you."
        ws_leads["A3"].font = Font(name="Arial", size=11, italic=True, color="27AE60")

    ws_leads.freeze_panes = "A2"

    # ----------------------------------------------------------------
    # SHEETS 3+: PER STAGE
    # ----------------------------------------------------------------
    enq_by_stage = defaultdict(list)
    for item in enq_alerts:
        enq_by_stage[item.get("stage", "Unknown")].append(item)

    e_hdrs = [
        "Customer Name", "Enquiry Name", "Enq No.",
        "kVA", "Offering", "Amount (₹)", "kVA Category", "Threshold",
        "Stalled For", "Enquiry Age (Days)", "Created Date",
        "Stage Since", "Remarks",
    ]

    stages = [st for st in config["stages_to_monitor"] if st not in config["terminal_stages"]]

    for stage in stages:
        display = stage_names.get(stage, stage)
        ws_stg = wb.create_sheet(display[:31])
        ws_stg.sheet_properties.tabColor = (
            config["stage_colors"].get(stage, "2c3e50").replace("#", "")
        )
        _hdr(ws_stg, e_hdrs, s)

        items = enq_by_stage.get(stage, [])
        if not items:
            ws_stg["A3"] = f"No stalled enquiries at {display} stage."
            ws_stg["A3"].font = Font(name="Arial", size=11, italic=True, color="27AE60")
            ws_stg.freeze_panes = "A2"
            continue

        enq_rows = []
        for item in items:
            rec = item["record"]
            is_na = item["threshold"] == 0
            stalled_for = "NA" if is_na else max(0, item["days_stalled"] - item["threshold"])
            threshold_display = "NA" if is_na else item["threshold"]
            age = _calc_age(rec)
            mod = get_field_value(rec, "Modified_Time", user_cache, crm)
            stage_since = mod[:10] if mod != "-" and len(mod) >= 10 else mod

            amt = None
            raw_amt = rec.get("Amount")
            if raw_amt and str(raw_amt) not in ("null", "None"):
                try:
                    amt = float(raw_amt)
                except (ValueError, TypeError):
                    pass

            enq_rows.append({
                "row": [
                    get_field_value(rec, "Account_Name", user_cache, crm),
                    get_field_value(rec, "Deal_Name", user_cache, crm),
                    get_field_value(rec, "ENQ_NUM", user_cache, crm),
                    get_field_value(rec, "DG_KVA", user_cache, crm),
                    get_field_value(rec, "OFFERING", user_cache, crm),
                    amt,
                    item["kva_category"],
                    threshold_display,
                    stalled_for,
                    age,
                    get_field_value(rec, "Created_Time", user_cache, crm)[:10],
                    stage_since,
                    get_field_value(rec, "Status_Remarks", user_cache, crm),
                ],
                "stalled_for": stalled_for if isinstance(stalled_for, int) else 0,
            })

        enq_rows.sort(key=lambda x: -x["stalled_for"])

        for ri, item in enumerate(enq_rows, 2):
            row = item["row"]
            for ci, val in enumerate(row, 1):
                c = ws_stg.cell(row=ri, column=ci, value=val)
                c.border = s["b"]
                c.font = s["df"]

            sf_val = item["stalled_for"]
            sf = ws_stg.cell(row=ri, column=9)
            if isinstance(sf_val, int):
                if sf_val >= 10:
                    sf.font = s["rf"]
                    sf.fill = s["redfill"]
                elif sf_val >= 5:
                    sf.font = s["bf"]
                    sf.fill = s["yellowfill"]

            if row[5] is not None:
                ws_stg.cell(row=ri, column=6).number_format = "#,##0"

        _autowidth(ws_stg, e_hdrs, len(enq_rows))
        _filter(ws_stg, e_hdrs, len(enq_rows))
        ws_stg.freeze_panes = "A2"

    # ----------------------------------------------------------------
    # SAVE
    # ----------------------------------------------------------------
    reports_dir = Path(__file__).parent.parent / "owner_reports"
    reports_dir.mkdir(exist_ok=True)
    safe_name = owner_name.replace(" ", "_").replace("/", "_")
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    filename = reports_dir / f"{safe_name}_{timestamp}.xlsx"
    wb.save(filename)
    return str(filename)
